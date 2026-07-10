
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"C:\Users\17873\Desktop\EduAgent")
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
DOCX_PATH = OUT / "EduAgent_前后端实现方案整理.docx"
MD_PATH = OUT / "EduAgent_前后端实现方案整理.md"


def read_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="utf-8", errors="ignore")
    except FileNotFoundError:
        return ""


def count_files_lines(folder: Path, suffixes=(".py", ".js", ".vue", ".css", ".json", ".md", ".yml", ".yaml", ".html")):
    files = 0
    lines = 0
    if not folder.exists():
        return files, lines
    skip_dirs = {"__pycache__", ".git", "node_modules", "dist", "venv", ".pytest_cache"}
    for p in folder.rglob("*"):
        if any(part in skip_dirs for part in p.parts):
            continue
        if p.is_file() and p.suffix.lower() in suffixes:
            files += 1
            try:
                lines += len(p.read_text(encoding="utf-8", errors="ignore").splitlines())
            except Exception:
                pass
    return files, lines


def extract_requirements():
    deps = []
    for line in read_text(ROOT / "requirements.txt").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        name = re.split(r"[<>=!~;\[]", line, maxsplit=1)[0].strip()
        deps.append((name, line))
    return deps


def extract_package_json():
    pkg_path = ROOT / "frontend" / "package.json"
    try:
        pkg = json.loads(read_text(pkg_path))
    except Exception:
        return {}, {}
    return pkg.get("dependencies", {}), pkg.get("devDependencies", {})


def extract_routes():
    text = read_text(ROOT / "frontend" / "server.py")
    routes = []
    pattern = re.compile(r'Route\("([^"]+)"\s*,\s*([a-zA-Z_][\w]*)\s*,\s*methods=\[([^\]]+)\]')
    for path, handler, methods in pattern.findall(text):
        method_list = ", ".join(re.findall(r'"([A-Z]+)"', methods))
        routes.append((method_list, path, handler))
    return routes


def extract_python_classes(path: Path):
    text = read_text(path)
    return re.findall(r"^class\s+([A-Za-z_][A-Za-z0-9_]*)", text, flags=re.M)


def extract_exported_functions(path: Path):
    text = read_text(path)
    funcs = re.findall(r"export\s+(?:async\s+)?function\s+([A-Za-z_][A-Za-z0-9_]*)", text)
    funcs += re.findall(r"export\s+const\s+([A-Za-z_][A-Za-z0-9_]*)", text)
    return funcs


src_files, src_lines = count_files_lines(ROOT / "src", suffixes=(".py", ".md", ".json", ".yml", ".yaml"))
fe_files, fe_lines = count_files_lines(ROOT / "frontend" / "src", suffixes=(".js", ".vue", ".css", ".json"))
test_files, test_lines = count_files_lines(ROOT / "tests", suffixes=(".py", ".md", ".json"))
requirements = extract_requirements()
fe_deps, fe_dev_deps = extract_package_json()
routes = extract_routes()
agent_state_classes = extract_python_classes(ROOT / "src" / "state" / "agent_state.py")
sql_classes = extract_python_classes(ROOT / "src" / "database" / "models_sqlalchemy.py")
edu_api_funcs = extract_exported_functions(ROOT / "frontend" / "src" / "services" / "eduAgentApi.js")


class Writer:
    def __init__(self):
        self.doc = Document()
        self.md: List[str] = []
        self._setup_doc()

    def _setup_doc(self):
        section = self.doc.sections[0]
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        styles = self.doc.styles
        for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
            if name in styles:
                style = styles[name]
                style.font.name = "Microsoft YaHei"
                style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles["Normal"].font.size = Pt(10.5)
        styles["Title"].font.size = Pt(24)
        for i, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
            if i in styles:
                styles[i].font.size = Pt(size)
                styles[i].font.bold = True
        try:
            code_style = styles.add_style("EduCode", 1)
            code_style.font.name = "Consolas"
            code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            code_style.font.size = Pt(8.5)
        except Exception:
            pass

    def title(self, text: str, subtitle: str | None = None):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(24)
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        self.md.append(f"# {text}\n")
        if subtitle:
            p2 = self.doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r2 = p2.add_run(subtitle)
            r2.font.size = Pt(12)
            r2.font.color.rgb = RGBColor(90, 90, 90)
            self.md.append(f"_{subtitle}_\n")

    def heading(self, text: str, level: int = 1):
        self.doc.add_heading(text, level=level)
        self.md.append("#" * (level + 1) + " " + text + "\n")

    def para(self, text: str = ""):
        p = self.doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.18
        if text:
            p.add_run(text)
        self.md.append(text + "\n")

    def bullets(self, items: Iterable[str], level: int = 0):
        for item in items:
            style = "List Bullet" if level == 0 else "List Bullet 2"
            self.doc.add_paragraph(item, style=style)
            self.md.append("  " * level + f"- {item}")
        self.md.append("")

    def numbered(self, items: Iterable[str]):
        for i, item in enumerate(items, 1):
            self.doc.add_paragraph(item, style="List Number")
            self.md.append(f"{i}. {item}")
        self.md.append("")

    def table(self, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float] | None = None):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        for i, h in enumerate(headers):
            self._set_cell_text(hdr[i], h, bold=True, shade="1F4E79", color="FFFFFF")
            if widths and i < len(widths):
                hdr[i].width = Cm(widths[i])
        for row in rows:
            cells = table.add_row().cells
            for i, val in enumerate(row):
                self._set_cell_text(cells[i], str(val))
                if widths and i < len(widths):
                    cells[i].width = Cm(widths[i])
        self.md.append("| " + " | ".join(headers) + " |")
        self.md.append("| " + " | ".join(["---"] * len(headers)) + " |")
        for row in rows:
            self.md.append("| " + " | ".join(str(x).replace("\n", "<br>") for x in row) + " |")
        self.md.append("")

    def _set_cell_text(self, cell, text, bold=False, shade=None, color=None):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(str(text))
        run.bold = bold
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(9)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if shade:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), shade)
            tcPr.append(shd)

    def code(self, text: str):
        p = self.doc.add_paragraph()
        try:
            p.style = "EduCode"
        except Exception:
            pass
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.right_indent = Cm(0.3)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8.5)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F2F2F2")
        pPr.append(shd)
        self.md.append("    " + text.replace("\n", "\n    ") + "\n")

    def page_break(self):
        self.doc.add_page_break()
        self.md.append("\n---\n")

    def save(self):
        self.doc.save(DOCX_PATH)
        MD_PATH.write_text("\n".join(self.md), encoding="utf-8")


w = Writer()

generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
w.title("EduAgent 项目前后端实现方案整理", f"生成时间：{generated_at}；项目路径：{ROOT}")
w.para("本文档基于当前工作区源码自动梳理并人工结构化整理，覆盖 EduAgent 项目的产品定位、总体架构、后端实现、前端实现、核心业务流程、API 契约、部署测试以及后续优化建议。文档重点面向技术评审、答辩汇报、团队交接与后续迭代开发。")
w.table(["项目", "内容"], [
    ["项目名称", "EduAgent：基于大模型的个性化资源生成与学习多智能体系统"],
    ["代码根目录", str(ROOT)],
    ["后端入口", "frontend/server.py，Starlette 应用，挂载 API 路由与前端静态资源"],
    ["核心后端包", "src：多智能体、应用服务、领域模型、数据库、图谱、向量检索、LLM、安全与观测"],
    ["前端入口", "frontend/src/main.js 与 frontend/src/App.vue，Vue 3 单页应用"],
    ["输出文档", str(DOCX_PATH)],
])
w.page_break()

w.heading("0. 文档目录", 1)
w.numbered([
    "项目总体概览",
    "总体架构与端到端业务链路",
    "后端实现方案",
    "前端实现方案",
    "核心功能清单与实现说明",
    "API 数据契约与接口设计",
    "部署、运行、观测与测试",
    "风险点、优化建议与迭代路线",
    "附录：关键文件地图",
])

w.heading("1. 项目总体概览", 1)
w.heading("1.1 项目定位", 2)
w.para("EduAgent 是一个面向学习者的智能教育系统。它以知识图谱为学习地图，以学习者画像为个性化依据，以多智能体协同为核心决策机制，动态完成冷启动画像、学习路径规划、知识点资源生成、Tutor 答疑、行为评估、掌握度更新、重规划与能力诊断。前端提供课程选择、画像问答、知识路径、资源画布、导师对话、能力雷达和诊断反馈等交互，后端负责会话状态、算法编排、资源生成、鉴权、持久化与观测。")
w.heading("1.2 代码规模", 2)
w.table(["模块", "文件数", "约代码行数", "说明"], [
    ["src", src_files, src_lines, "后端核心包：Agent、应用服务、数据库、图谱、向量、LLM、领域模型等"],
    ["frontend/src", fe_files, fe_lines, "Vue 3 前端源码：页面、组件、组合式状态、API 客户端、样式"],
    ["tests", test_files, test_lines, "单元测试与集成测试，覆盖算法节点、应用服务、鉴权和 HTTP 行为"],
])
w.heading("1.3 技术栈总览", 2)
w.table(["层级", "技术", "用途"], [
    ["前端框架", "Vue 3、Vue Router、Vite", "构建单页学习工作台、路由和开发构建工具链"],
    ["前端交互", "Axios、Marked、DOMPurify、Mermaid", "HTTP 通信、Markdown 渲染、安全净化、图表渲染"],
    ["前端样式", "TailwindCSS、PostCSS、Autoprefixer", "响应式布局、设计系统和浏览器兼容"],
    ["HTTP 服务", "Starlette、FastAPI 依赖、Uvicorn、SSE", "API 路由、静态资源托管、流式输出"],
    ["多智能体编排", "LangGraph、LangChain Core、自研 orchestration_core", "评估、画像、规划、资源、校验、Tutor、评测协同"],
    ["数据模型", "Pydantic v2、SQLAlchemy asyncio", "DTO 校验、AgentState 校验、异步数据库模型"],
    ["知识图谱", "Neo4j", "知识点、先修关系、路径规划、图谱查询"],
    ["检索与向量", "Milvus、Elasticsearch、sentence-transformers", "父子分块、向量检索、关键词检索、资源增强"],
    ["LLM 接入", "OpenAI SDK、HTTPX、自研 client_v2", "DashScope、DeepSeek、OpenAI、Spark 等 Provider 适配与降级"],
    ["安全鉴权", "argon2-cffi、JWT、python-jose、passlib、验证码、限流", "注册登录、Token 刷新、密码哈希、接口保护"],
    ["可观测性", "loguru、自研 metrics/request context", "请求 ID、结构化日志、指标快照、耗时统计"],
    ["部署", "Docker、docker-compose、Nginx 目录", "前端构建、Python 运行时、Neo4j/ES/Milvus/MinIO/etcd 编排"],
])
w.heading("1.4 主要依赖清单", 2)
backend_deps_rows = []
for name, line in requirements[:28]:
    backend_deps_rows.append([name, line])
w.table(["后端依赖", "版本约束或声明"], backend_deps_rows)
frontend_rows = [[k, v, "runtime"] for k, v in fe_deps.items()] + [[k, v, "dev"] for k, v in fe_dev_deps.items()]
w.table(["前端依赖", "版本", "类型"], frontend_rows)

w.heading("2. 总体架构与端到端业务链路", 1)
w.heading("2.1 分层架构", 2)
w.para("项目采用前后端分离但同容器交付的模式：Vite 构建后的前端静态资源由 Starlette 服务托管；同一个 Starlette 应用暴露业务 API、鉴权 API、兼容 API 与观测 API。后端内部按入口层、应用服务层、编排层、智能体层、领域状态层、数据基础设施层、外部模型/检索层分层。")
w.code("用户浏览器\n  -> Vue 3 SPA：路由、工作台、资源画布、Tutor 对话、知识树\n  -> Axios API Client：Bearer Token、刷新 Token、X-Request-ID\n  -> Starlette frontend/server.py：API 路由、SSE、静态资源\n  -> Application Services：Session/Profile/Resource/Tutor/Assessment\n  -> OrchestrationRuntime + run_official_learning_step\n  -> Multi-Agent Nodes：Evaluator、Profiler、Planner、Tutor、ContentMesh、Validator、Assessment\n  -> Domain State：AgentState、LearningSession、ResourceCard、DynamicProfile\n  -> Persistence & Infra：PostgreSQL/SQLite、Neo4j、Milvus、Elasticsearch、Redis、MinIO\n  -> LLM Providers：DashScope、DeepSeek、OpenAI、Spark")
w.heading("2.2 核心闭环", 2)
w.table(["阶段", "前端表现", "后端处理", "输出结果"], [
    ["登录与身份", "AuthView 展示登录/注册/验证码", "auth 路由校验验证码、哈希密码、签发 access/refresh token", "用户身份与课程档案"],
    ["课程选择", "CourseSelectionView 展示可选课程和已选课程", "CourseRepo/EnrollmentRepo 获取课程、选课和切课", "activeCourse、courseId、知识图谱入口"],
    ["冷启动画像", "ChatArea 逐题采集画像答案", "ColdStartOrchestrator/ProfileService 更新 StaticProfile", "认知风格、动机、基础水平、目标节点"],
    ["路径初始化", "知识树高亮 activePath", "PlannerNode 基于图谱、掌握度、目标节点规划路径", "有序知识点路径、当前节点"],
    ["节点学习", "ResourceCanvas 渲染概念图、代码、练习、测验等卡片", "ContentMesh/ResourceService 生成或读取资源，并由 Validator 校验", "按认知风格适配的 ResourceCard"],
    ["Tutor 答疑", "ChatArea 支持普通与流式回答", "TutorAgentNode 结合当前节点、父上下文和画像生成解释", "个性化答疑、引导式解释"],
    ["行为评估", "用户提交正确率、耗时、求助、反馈", "Evaluator 清洗行为、PID 判断重规划、Profiler 更新掌握度", "mastery、replan flag、错误分布"],
    ["诊断评测", "能力雷达和诊断报告展示", "AssessmentReporterNode 计算 5 维能力、EMA 与策略", "diagnostic_report_md、capability_radar"],
])
w.heading("2.3 用户主流程", 2)
w.numbered([
    "访问首页，进入应用工作台；若无有效登录态则进入登录/注册流程。",
    "登录成功后拉取用户信息、已选课程、可选课程；没有当前课程时进入课程选择。",
    "创建或恢复学习 session，若画像未完成则进入 profile probe 冷启动问答。",
    "画像完成后初始化学习路径，前端进入 ready 模式并加载当前知识点资源。",
    "用户阅读资源、完成练习、向 Tutor 提问，前端将行为数据提交给后端。",
    "后端运行官方学习步骤：Evaluator、Profiler、Planner、Tutor、ContentMesh、Validator、Assessment 协同更新状态。",
    "前端接收 SessionResponse，刷新路径、资源、Tutor 消息、Agent 时间线、能力雷达和诊断报告。",
    "当掌握度达到阈值时推进到下一个知识点；当 PID 或策略触发时重规划路径或调整教学策略。",
])

w.heading("3. 后端实现方案", 1)
w.heading("3.1 后端目录分层", 2)
w.table(["目录或文件", "职责"], [
    ["src/application", "应用服务层，封装 session、profile、resource、tutor、assessment 等用例，屏蔽 HTTP 入口与编排细节"],
    ["src/agents", "多智能体节点与资源生成 Agent，包括评估、画像、规划、资源网络、校验、导师、评测等"],
    ["src/state/agent_state.py", "核心状态模型，定义 StaticProfile、DynamicProfile、LatestBehavior、ResourceCard、AgentState 等"],
    ["src/api_models", "面向 HTTP 返回的响应模型，包括 session/resource/profile/tutor DTO"],
    ["src/domain", "领域实体、课程、学习会话、知识节点、资源等领域模型"],
    ["src/database", "SQLAlchemy 异步模型与 repository，负责用户、课程、会话、资源、评测、快照等持久化"],
    ["src/auth", "验证码、JWT、中间件、速率限制、提示注入防御、认证路由"],
    ["src/graph", "Neo4j 客户端和知识图谱管理器，提供知识节点、边、邻接、课程图谱等查询"],
    ["src/vector", "Milvus 与 Elasticsearch 相关检索实现，支撑资源内容增强"],
    ["src/llm", "LLM Provider 客户端、输入管理、Token 估计、内容过滤、幻觉检测"],
    ["src/infrastructure", "冷启动、文档分块、图谱构建、传统路径规划、PID 控制等基础算法"],
    ["src/observability.py", "请求上下文、request_id、结构化日志、内存指标、耗时统计"],
    ["src/orchestration_core.py", "官方多智能体学习步骤和 LangGraph 兼容编排"],
    ["src/orchestration_runtime.py", "运行时单例，集中创建 KG、LLM、各 Agent 节点和 session runtime"],
    ["frontend/server.py", "实际 Web 入口，集中注册 API 路由、兼容路由、鉴权路由和前端静态资源"],
])
w.heading("3.2 HTTP 入口与路由", 2)
w.para("当前服务入口位于 frontend/server.py。它不是单纯的前端目录文件，而是生产运行时的 Starlette 应用入口：注册 REST API、SSE 流式接口、鉴权接口、旧版兼容接口，并在构建产物存在时托管 frontend/dist 静态文件。Dockerfile 最终命令也是 python frontend/server.py。")
route_rows = [[m, p, h] for m, p, h in routes]
w.table(["方法", "路径", "处理函数"], route_rows)
w.heading("3.3 应用服务层", 2)
w.table(["服务", "主要职责", "典型调用场景"], [
    ["SessionService", "创建、恢复、推进学习会话；封装 LearningSession 与 AgentState 同步；生成 SessionResponse", "createSession、getSession、advance、behavior、replan"],
    ["ProfileService", "画像探针、画像答案提交、冷启动完成判断、静态画像写入", "profile-probe、profile-input、cold-start 兼容接口"],
    ["ResourceService", "按 node_id 读取或生成资源；统一资源契约；接入 ContentMesh 与 Validator", "打开节点、切换节点、资源刷新"],
    ["TutorService", "普通 Tutor 与 SSE 流式 Tutor；组合上下文、调用 TutorAgentNode", "tutor、tutor-stream、legacy ask"],
    ["AssessmentService", "生成诊断反馈、能力雷达、评测报告；对接 AssessmentReporterNode", "节点推进后评估、诊断报告刷新"],
    ["_common", "通用会话解析、错误处理、DTO 组装、响应转换辅助", "各服务共用"],
])
w.heading("3.4 官方学习步骤编排", 2)
w.para("run_official_learning_step 是当前后端最关键的学习闭环入口。它将一次用户交互转化为可追踪的 Agent 流水线，确保状态更新、资源生成、Tutor 和评测在统一语义下执行。")
w.numbered([
    "解析输入：确定 current_node、correctness、time_spent_ratio、code_pass_rate、help_count、tutor_query 与 interaction_type。",
    "Evaluator：在非 load_node 场景下清洗行为向量，识别异常和错误类型，计算 PID 控制量并决定是否触发重规划。",
    "Profiler：更新动态画像，包括知识掌握度、错误分布、连续失败计数、遗忘曲线相关记录。",
    "Planner：当 active_path 为空或重规划标记为真时，结合知识图谱和掌握度重新计算学习路径。",
    "Tutor：当存在 tutor_query 且不是单纯加载节点时，调用 TutorAgentNode 生成个性化答疑。",
    "ContentMesh/ResourceService：为当前节点生成或读取资源卡片，并处理预生成与缓存。",
    "Validator：对资源进行符号/AST 与 NLI 层面的反幻觉校验，校验结果写入资源元数据和 Agent 反馈。",
    "Assessment：生成能力雷达、策略控制和诊断报告，更新 dynamic_profile。",
    "推进节点：在诊断或测验达到掌握阈值后切换 current_node，维持 active_path 状态。",
    "返回 LearningStepResult：最终由应用服务转换为前端需要的 SessionResponse。",
])
w.code("run_official_learning_step(input, runtime_session)\n  -> evaluator_output\n  -> profiler_output / dynamic_profile\n  -> planner_output / active_path\n  -> tutor_response optional\n  -> resource_cards + validator status\n  -> assessment_report\n  -> updated AgentState + LearningStepResult")
w.heading("3.5 OrchestrationRuntime 运行时", 2)
w.para("OrchestrationRuntime 是稳定的官方运行时单例，集中管理 LLM Provider、知识图谱管理器、各类 Agent 节点和 session runtime。这样可以避免 HTTP 请求中反复构造模型客户端和图谱连接，也使测试可以替换 session state 或重置 runtime。")
w.table(["成员", "说明"], [
    ["provider order", "从 LLM_PROVIDER 和 EDUAGENT_LLM_FALLBACKS 构建，例如 deepseek、openai、dashscope、spark；只在凭据存在时启用"],
    ["kg", "get_kg_manager 创建的知识图谱管理器，向 Planner 和课程图谱 API 提供数据"],
    ["EvaluatorNode", "行为评估、PID 重规划、错误分布更新"],
    ["ProfilerNode", "认知画像与掌握度更新，包含 Thompson Sampling 与遗忘曲线思想"],
    ["PlannerNode", "图谱路径规划，使用 DAG-Dijkstra 和拓扑约束"],
    ["TutorAgentNode", "个性化答疑，接入 LLM Generator"],
    ["ContentMeshNode", "资源调度，支持 WFQ 队列和 Markov shadow 预生成"],
    ["ValidatorNode", "资源校验，结合符号规则、AST、NLI"],
    ["AssessmentReporterNode", "能力雷达、EMA、诊断报告和策略控制"],
])
w.heading("3.6 多智能体节点设计", 2)
w.table(["Agent/节点", "输入", "核心算法或策略", "输出/副作用"], [
    ["EvaluatorNode", "LatestBehavior、当前节点、历史掌握度", "行为清洗、异常检测、错误类型分类、PIDReplanController", "EvaluatorOutput、re_plan_triggered、错误分布、Agent 反馈"],
    ["ProfilerNode", "Evaluator 结果、资源反馈、历史画像", "Thompson Sampling 调整资源偏好；Ebbinghaus 遗忘曲线修正掌握度", "DynamicProfile、recommended_resource_style"],
    ["PlannerNode", "知识图谱、target_node、mastery、约束", "DAG-Dijkstra、拓扑排序、先修约束、边权调整", "active_path、当前节点建议、路径边信息"],
    ["ContentMeshNode", "node_id、画像、路径、资源需求", "Weighted Fair Queuing、Markov shadow pre-generation、缓存命中", "ResourceCard 列表、任务状态、预生成缓存"],
    ["ValidatorNode", "资源内容、知识点、上下文", "Bipolar anti-hallucination：符号/AST 门禁 + 滑动窗口 NLI", "校验状态、风险标签、修正建议"],
    ["TutorAgentNode", "tutor_query、当前节点、父上下文、画像", "苏格拉底式/讲解式/提示式答疑、Mermaid 语法防护、上下文裁剪", "TutorResponse 或 SSE chunk"],
    ["AssessmentReporterNode", "掌握度、行为、错误、能力历史", "EMA 能力雷达、边界遗漏趋势、Hysteresis 策略控制", "5 维能力雷达、diagnostic_report_md、教学策略"],
    ["ResourceGenerationAgent", "资源类型、知识点、目标难度、上下文", "统一资源生成入口，分发 PPT、quiz、mindmap、coding、video 等生成策略", "规范化资源内容和 metadata"],
])
w.heading("3.7 AgentState 与核心领域状态", 2)
w.para("AgentState 是多智能体协作的共享状态。它将用户、课程、静态画像、动态画像、最新行为、学习路径、资源、反馈和诊断结果组织在同一个强类型 Pydantic 模型中，便于在 Agent 节点之间安全传递并最终序列化给前端。")
w.table(["模型", "关键字段", "用途"], [
    ["StaticProfile", "cognitive_style、motivation、time_budget、knowledge_base", "冷启动后形成的长期静态画像，影响资源样式和路径难度"],
    ["DynamicProfile", "knowledge_mastery、records、error_type_distribution、continuous_fail_counter、pid_errors、capability_radar、diagnostic_report_md", "动态学习状态，随每轮交互持续更新"],
    ["LatestBehavior", "node_id、correctness、time_spent_ratio、resource_feedback、help_request_count、tutor_query、accuracy_rate、code_pass_rate、duration_ratio", "前端提交的本轮学习行为向量"],
    ["ResourceCard", "resource_id、node_id、card_type、title/content/status/metadata", "前端 ResourceCanvas 渲染的标准资源卡片契约"],
    ["AgentFeedbackItem", "agent、stage、status、headline、summary、details_md、structured_data、artifacts", "多智能体执行过程向前端解释的时间线"],
    ["AgentState", "user_id、course_id、current_node_id、target_node_id、static_profile、dynamic_profile、active_path、generated_resources、latest_behavior", "官方学习闭环的完整状态载体"],
])
w.para("当前 agent_state.py 中识别到的主要类包括：" + "、".join(agent_state_classes[:30]) + "。")
w.heading("3.8 数据持久化与仓库层", 2)
w.para("数据库层使用 SQLAlchemy asyncio 建模，按 Repository 拆分用户、课程、选课、会话、状态、资源、评测等数据访问。开发或测试场景可使用 SQLite/fakeredis，生产 compose 中预留 PostgreSQL、Redis 等服务。")
w.table(["Repository/模型", "职责说明"], [
    ["user_repo、User", "用户注册、登录查询、用户资料"],
    ["course_repo、KnowledgePoint", "课程列表、课程详情、知识点基础信息"],
    ["enrollment_repo、user_course_profile_repo", "用户选课、切换当前课程、课程级画像"],
    ["session_repo、LearningSessionModel", "学习 session 的创建、状态、当前节点、目标节点"],
    ["state_repo、session_snapshot_repo", "AgentState 快照保存、会话恢复、回放分析"],
    ["resource_repo、GeneratedResource", "资源卡片缓存、生成结果、校验 metadata"],
    ["evaluation_repo、EvaluationReport", "诊断报告、能力雷达、测验记录"],
    ["redis_client", "缓存、限流或短期状态存储能力"],
])
w.para("models_sqlalchemy.py 中识别到的类包括：" + "、".join(sql_classes[:40]) + "。")
w.heading("3.9 知识图谱、检索增强与资源生成", 2)
w.bullets([
    "知识图谱：src/graph/knowledge_graph_manager.py 和 neo4j_client.py 负责访问 Neo4j，向课程图谱 API、PlannerNode 与父上下文检索提供节点和边。",
    "路径规划：PlannerNode 基于图谱先修关系构建 DAG，结合掌握度和目标节点选择最适合当前学习者的路径。",
    "文档分块：document_chunker.py 提供父子块或资源文档切分能力，为向量检索提供粒度基础。",
    "向量与全文检索：Milvus 适合语义相似检索，Elasticsearch 适合关键词/结构化检索，两者共同服务资源增强和上下文召回。",
    "资源生成：ResourceGenerationAgent 将概念图、代码片段、互动练习、测验、思维导图、视频脚本等统一成 ResourceCard。",
    "资源校验：ValidatorNode 在生成后对事实一致性、代码语法/AST、内容边界和 NLI 置信度做检查，降低幻觉风险。",
])
w.heading("3.10 LLM 接入、降级与安全防护", 2)
w.table(["能力", "实现要点"], [
    ["Provider 适配", "src/llm/client_v2.py 与 Runtime 根据环境变量选择 dashscope、spark、deepseek、openai 等 Provider"],
    ["降级策略", "EDUAGENT_LLM_FALLBACKS 指定回退顺序；缺少凭据的 Provider 不会进入可用池"],
    ["输入管理", "input_manager、token_estimator 对上下文进行裁剪、估算和格式化，避免超过模型窗口"],
    ["内容过滤", "content_filter 和 prompt_defense 处理不安全输入、提示注入或越权请求"],
    ["幻觉检测", "hallucination_checker 与 ValidatorNode 的 NLI/符号校验形成双层防线"],
    ["Tutor 防护", "TutorAgentNode 对 Mermaid、代码块、教学模式和父上下文进行限制，避免前端渲染异常和无关输出"],
])
w.heading("3.11 鉴权、会话安全与可观测性", 2)
w.bullets([
    "鉴权流程：验证码生成与校验、注册登录、argon2/passlib 密码哈希、JWT access token 与 refresh token。",
    "前端携带 Bearer token 和 withCredentials，后端 /api/auth/refresh 支持续期，失败时前端清空登录态并跳转。",
    "rate_limiter 可用于限制登录、验证码或高成本接口调用，降低暴力破解和滥用风险。",
    "observability.py 提供 new_request_id、bind_context、log_event、incr_metric、observe_metric、timed_operation 等能力。",
    "每个请求通过 X-Request-ID 串联前端、后端日志和指标，/api/ops/metrics 提供运行期指标快照。",
])

w.heading("4. 前端实现方案", 1)
w.heading("4.1 前端总体结构", 2)
w.table(["目录/文件", "职责"], [
    ["frontend/src/main.js", "创建 Vue 应用，挂载路由和全局样式"],
    ["frontend/src/App.vue", "应用根组件，承载 router-view"],
    ["frontend/src/router/index.js", "路由定义，/ 为 LandingView，/app 为 AppWorkspaceView，其他重定向"],
    ["frontend/src/services/apiClient.js", "Axios 实例、请求 ID、Token 注入、401 刷新、错误处理"],
    ["frontend/src/services/eduAgentApi.js", "业务 API 封装，隐藏 REST 路径细节"],
    ["frontend/src/composables/useEduAgent.js", "核心组合式状态机，统一管理登录、课程、画像、路径、资源、聊天、诊断"],
    ["frontend/src/views/AppWorkspaceView.vue", "应用工作台容器，根据 bootMode 切换登录、选课、画像、ready 状态"],
    ["frontend/src/components", "Landing、Auth、CourseSelection、PremiumWorkspace、ResourceCanvas、ChatArea、SidebarDrawer 等 UI 组件"],
])
w.heading("4.2 路由与启动状态", 2)
w.table(["路由/状态", "说明"], [
    ["/", "LandingView：产品首页/品牌页，引导进入应用"],
    ["/app", "AppWorkspaceView：核心学习工作台"],
    ["fallback", "未知路径重定向到 /"],
    ["bootMode=loading", "应用初始化、检查登录态、加载用户和课程"],
    ["bootMode=login", "未登录或刷新失败，显示 AuthView"],
    ["bootMode=course_selection", "已登录但未选择课程，展示课程选择"],
    ["bootMode=probe", "会话存在但画像未完成，进入画像问答"],
    ["bootMode=ready", "课程、画像、路径、当前节点资源均可用，进入学习工作台"],
])
w.heading("4.3 useEduAgent 组合式状态", 2)
w.para("useEduAgent.js 是前端事实上的状态中枢。它将 API 调用、业务状态、计算属性、用户动作和界面反馈组织为一个组合式 hook，供 AppWorkspaceView 和子组件使用。")
w.table(["状态类别", "字段示例", "用途"], [
    ["身份状态", "isLoggedIn、currentUser、userId", "控制登录态、用户显示、Token 刷新后的恢复"],
    ["课程状态", "activeCourse、availableCourses、enrolledCourses、courseId", "课程选择、切课、课程图谱加载"],
    ["启动状态", "bootMode、isBusy、isSubmittingProbe、isLoadingNode", "控制页面阶段和加载按钮"],
    ["学习状态", "currentNode、activePath、mastery、knowledgeGraph、nodeTitles", "知识树、路径进度、当前知识点标题"],
    ["画像状态", "probe、probeCollected、probeTotal", "冷启动问答进度"],
    ["资源状态", "resources、currentCards、lastDiagnostic", "ResourceCanvas 渲染和节点内容加载"],
    ["Agent 反馈", "agentFeedback、stepLogs、agentStatuses", "展示多智能体阶段、状态、摘要"],
    ["聊天状态", "messages、infoMessage", "Tutor 问答、画像对话、系统提示"],
    ["评测状态", "capabilityRadar、diagnosticReport、overallProgress、masteredCount", "能力雷达、诊断报告和整体进度"],
])
w.heading("4.4 API 通信与 Token 刷新", 2)
w.bullets([
    "apiClient.js 创建 baseURL 为 /api 的 Axios 实例，默认超时 30000 ms，withCredentials=true。",
    "每个请求自动添加 X-Request-ID，便于后端日志串联。",
    "若存在 access token，则自动加入 Authorization: Bearer token。",
    "收到 401 后自动调用 /api/auth/refresh；刷新成功后重放原请求，刷新失败则清空本地 token 并回到登录态。",
    "eduAgentApi.js 只暴露业务语义函数，组件不直接拼接 URL，从而降低接口变更成本。",
])
w.para("eduAgentApi.js 中识别到的导出函数包括：" + "、".join(edu_api_funcs) + "。")
w.heading("4.5 主要组件职责", 2)
w.table(["组件", "职责", "关键交互"], [
    ["LandingView", "首页与产品介绍", "进入应用、展示品牌和价值主张"],
    ["AuthView", "登录、注册、验证码", "调用 captcha/register/login，成功后通知父层启动工作台"],
    ["CourseSelectionView", "课程列表、已选课程、选课/切课", "enrollCourse、switchCourse、fetchCourseDetail"],
    ["AppWorkspaceView", "工作台总容器与状态分发", "根据 bootMode 渲染不同阶段，连接 useEduAgent"],
    ["PremiumWorkspace", "学习空间主布局", "承载侧边栏、资源区、聊天区、进度/诊断模块"],
    ["ResourceCanvas", "资源卡片渲染", "根据 card_type 渲染概念、代码、练习、测验、图表等"],
    ["ChatArea", "Tutor 对话和画像问答", "提交用户问题、显示流式回答、采集探针答案"],
    ["SidebarDrawer", "侧边导航与学习设置", "展示路径、节点状态、课程信息、设置入口"],
    ["KnowledgeTree", "知识图谱/路径可视化", "高亮 currentNode、activePath、mastery 状态"],
    ["MarkdownContent", "安全 Markdown 渲染", "Marked 转 HTML，DOMPurify 净化，Mermaid 渲染图表"],
    ["RadarCanvas", "能力雷达图", "展示 AssessmentReporterNode 生成的 5 维能力"],
])
w.heading("4.6 前端资源契约", 2)
w.para("前端围绕 ResourceCard 进行渲染。后端必须保证每张卡片具备稳定的 resource_id、node_id、card_type、title/content/status/metadata。content 通常为 Markdown，metadata 承载难度、校验状态、渲染类型、题目答案、代码语言等扩展字段。")
w.table(["card_type", "前端渲染倾向", "典型内容"], [
    ["concept_map", "Markdown + Mermaid/图谱", "概念关系、先修关系、关键定义"],
    ["code_snippet", "代码块与解释", "示例代码、复杂度说明、运行提示"],
    ["interactive_exercise", "练习卡片", "填空、选择、操作步骤、反馈"],
    ["quiz", "测验卡片", "题干、选项、答案、解析、掌握度提交"],
    ["mindmap", "Mermaid 或层级结构", "知识点树状梳理"],
    ["video_summary", "脚本/摘要", "视频讲解脚本、知识点时间线"],
])

w.heading("5. 核心功能清单与实现说明", 1)
w.table(["功能域", "功能", "后端实现", "前端实现"], [
    ["用户", "注册/登录/刷新/退出/个人信息", "src/auth + /api/auth/*", "AuthView + apiClient token 管理"],
    ["课程", "课程列表、详情、选课、切课", "course_repo、enrollment_repo、/api/courses、/api/user/courses", "CourseSelectionView、useEduAgent 课程状态"],
    ["画像", "冷启动探针、画像答案、静态画像", "ColdStartOrchestrator、ProfileService", "probe bootMode、ChatArea 问答"],
    ["路径", "初始化路径、重规划、知识树", "PlannerNode、KnowledgeGraphManager", "KnowledgeTree、activePath、currentPathNodes"],
    ["学习", "节点加载、行为提交、掌握度更新", "run_official_learning_step、Evaluator、Profiler", "ResourceCanvas、submitSessionBehavior、advanceSession"],
    ["资源", "个性化资源生成、缓存、校验", "ContentMesh、ResourceGenerationAgent、Validator、ResourceRepo", "currentCards、MarkdownContent"],
    ["答疑", "普通 Tutor、SSE 流式 Tutor", "TutorService、TutorAgentNode、tutor-stream", "ChatArea、messages、streamSessionTutor"],
    ["诊断", "能力雷达、诊断报告、策略控制", "AssessmentReporterNode、AssessmentService", "RadarCanvas、diagnosticReport"],
    ["观测", "日志、请求 ID、指标", "observability.py、/api/ops/metrics", "X-Request-ID、Agent 时间线展示"],
    ["兼容", "旧版 API 兼容", "/api/state、/api/pipeline/*、/api/tutor/* 等", "便于旧前端或测试脚本继续调用"],
])
w.heading("5.1 冷启动画像流程", 2)
w.numbered([
    "前端创建 session 后调用 profile-probe 获取下一道画像问题。",
    "用户在 ChatArea 中作答，前端调用 profile-input 提交答案。",
    "ProfileService 将答案交给 ColdStartOrchestrator，更新 StaticProfile 的认知风格、动机、基础知识、时间预算等字段。",
    "当 is_complete 为真时 finalize 画像，确定 target_node 或初始学习目标。",
    "前端从 probe 模式切换到路径初始化，随后进入 ready。",
])
w.heading("5.2 节点学习与资源生成流程", 2)
w.numbered([
    "前端根据 currentNode 请求 /api/sessions/{session_id}/resources/{node_id} 或在 advance/behavior 响应中获取资源。",
    "ResourceService 检查资源缓存；缺失或过期时调用 ContentMeshNode 创建生成任务。",
    "ContentMeshNode 使用 WFQ 调度不同资源类型，保障测验、概念解释、代码练习等资源公平生成。",
    "ResourceGenerationAgent 根据课程、节点、画像和资源类型生成 Markdown/结构化内容。",
    "ValidatorNode 对内容进行规则校验、AST 校验或 NLI 校验，写入 status 和 metadata。",
    "前端 ResourceCanvas 按 card_type 渲染，MarkdownContent 对 HTML 进行 DOMPurify 净化并渲染 Mermaid。",
])
w.heading("5.3 测验推进与重规划闭环", 2)
w.numbered([
    "用户完成测验或练习后，前端整理 correctness、accuracy_rate、code_pass_rate、duration_ratio、help_request_count 等行为数据。",
    "EvaluatorNode 清洗行为并识别异常，例如耗时异常、连续错误、求助过多、掌握度偏离。",
    "PIDReplanController 根据误差累计和变化趋势判断是否需要重规划。",
    "ProfilerNode 更新对应 node_id 的 knowledge_mastery，并记录错误类型分布和能力信号。",
    "如果掌握度达到阈值，SessionService 推进到 activePath 的下一个节点；如果触发重规划，则 PlannerNode 重新计算路径。",
    "AssessmentReporterNode 更新能力雷达和诊断报告，前端同步刷新进度、知识树和建议。",
])

w.heading("6. API 数据契约与接口设计", 1)
w.heading("6.1 REST API 分类", 2)
w.table(["类别", "接口", "说明"], [
    ["Session", "/api/sessions、/api/sessions/{session_id}", "创建与读取学习会话"],
    ["Profile", "/api/sessions/{session_id}/profile-probe、profile-input", "冷启动画像问题与答案"],
    ["Path", "/api/sessions/{session_id}/path/init、advance、behavior、replan", "路径初始化、推进、行为提交、重规划"],
    ["Tutor", "/api/sessions/{session_id}/tutor、tutor-stream", "普通答疑与流式答疑"],
    ["Resource", "/api/sessions/{session_id}/resources/{node_id}", "获取某节点资源卡片"],
    ["Course", "/api/courses、/api/courses/{course_id}", "课程列表与详情"],
    ["User Course", "/api/user/courses、enroll、switch", "用户课程关系"],
    ["Graph", "/api/knowledge-graph", "知识图谱数据"],
    ["Auth", "/api/auth/captcha、register、login、refresh、logout、me", "验证码与身份认证"],
    ["Ops", "/api/ops/metrics、/api/reset", "指标和调试重置"],
])
w.heading("6.2 SessionResponse 结构", 2)
w.para("SessionResponse 是前端最核心的响应契约，目标是让一次后端调用即可刷新工作台所需大部分状态。它通常包含 session、profile、dynamic_profile、learning_path、resources、assessment、tutor_response、agent_feedback、pipeline_log 等字段，并带有 resource_contract_version 与 agent_feedback_version 以支持前后端协同演进。")
w.table(["字段", "含义", "前端消费方"], [
    ["session", "学习会话元数据：user_id、course_id、current_node_id、target_node_id、iteration、status", "useEduAgent、PremiumWorkspace"],
    ["profile", "静态画像：认知风格、动机、基础、时间预算", "资源样式、画像完成判断"],
    ["dynamic_profile", "掌握度、错误分布、能力雷达、诊断报告", "KnowledgeTree、RadarCanvas、诊断面板"],
    ["learning_path", "active_path 和节点信息", "知识树、路径进度、下一节点"],
    ["resources", "按 node_id 或当前节点聚合的 ResourceCard 列表", "ResourceCanvas"],
    ["assessment", "评测输出、能力分数、策略建议", "诊断报告、雷达图"],
    ["tutor_response", "Tutor 普通回答或最终汇总", "ChatArea"],
    ["agent_feedback", "Agent 时间线，包含每个阶段状态和结构化数据", "Agent 状态条、调试面板"],
    ["pipeline_log", "流水线日志或兼容模式日志", "开发调试、演示说明"],
])
w.heading("6.3 关键请求字段", 2)
w.table(["请求", "关键字段", "说明"], [
    ["createSession", "user_id、course_id、target_node_id 可选", "创建或恢复用户在课程下的学习会话"],
    ["profile-input", "answer、question_id、metadata", "提交画像探针答案"],
    ["path/init", "target_node_id、force", "初始化学习路径"],
    ["advance/behavior", "node_id、correctness、time_spent_ratio、resource_feedback、help_request_count、accuracy_rate、code_pass_rate、duration_ratio", "提交学习行为并推进闭环"],
    ["tutor", "query、node_id、mode、context", "发起导师问答"],
    ["replan", "reason、target_node_id 可选", "用户或系统触发重规划"],
])
w.heading("6.4 SSE 流式接口", 2)
w.para("Tutor 流式接口适合长回答或逐步解释。后端使用 sse-starlette/Starlette 流式能力分片输出，前端 streamSessionTutor 将 chunk 追加到当前 assistant 消息。实现上需要保证异常 chunk、结束事件、Token 过期重试和用户取消的处理。")

w.heading("7. 部署、运行、观测与测试", 1)
w.heading("7.1 Docker 部署", 2)
w.para("Dockerfile 采用多阶段构建：第一阶段使用 Node 构建 Vue 前端，第二阶段使用 python:3.11-slim 安装 requirements，复制 src、tests、scripts、frontend/server.py 和前端 dist，暴露 8800 端口，健康检查调用 /api/state，启动命令为 python frontend/server.py。")
w.table(["docker-compose 服务", "用途"], [
    ["app", "EduAgent Web/API 应用容器"],
    ["neo4j", "知识图谱数据库"],
    ["elasticsearch", "全文检索服务"],
    ["etcd", "Milvus 依赖的元数据服务"],
    ["minio", "Milvus 对象存储依赖，也可扩展为资源文件存储"],
    ["milvus-standalone", "向量数据库服务"],
])
w.heading("7.2 环境变量", 2)
w.table(["类别", "变量示例", "说明"], [
    ["LLM", "LLM_PROVIDER、DASHSCOPE_API_KEY、SPARK_API_KEY、DEEPSEEK_API_KEY、OPENAI_API_KEY", "选择模型 Provider 和 API 密钥"],
    ["数据库", "DATABASE_URL、REDIS_URL", "关系数据库与缓存连接"],
    ["图谱", "NEO4J_URI、NEO4J_USER、NEO4J_PASSWORD", "Neo4j 连接"],
    ["检索", "ES_HOSTS、ES_USER、ES_PASSWORD、CHROMA_PERSIST_DIR", "全文检索或迁移期向量库配置"],
    ["对象存储", "MINIO_ACCESS_KEY、MINIO_SECRET_KEY", "Milvus/资源对象存储配置"],
    ["安全", "JWT_SECRET_KEY", "JWT 签名密钥，生产必须替换"],
    ["模型下载", "HF_ENDPOINT", "国内镜像或模型下载端点"],
])
w.heading("7.3 测试覆盖", 2)
w.para("tests 目录覆盖了核心算法、Agent 节点、应用服务、鉴权、HTTP 可观测性和前端资源契约。建议在提交前执行 pytest，并在需要时补充端到端用例。")
w.table(["测试文件/类别", "覆盖内容"], [
    ["test_pid_controller.py", "PID 控制和重规划判断"],
    ["test_path_planner.py、test_planner_node.py", "路径规划、DAG、拓扑和节点输出"],
    ["test_profiler_node.py", "画像与掌握度更新"],
    ["test_evaluator_node.py", "行为清洗、异常检测、错误分布"],
    ["test_tutor_node.py", "Tutor 节点输出和上下文行为"],
    ["test_cold_start.py", "冷启动画像流程"],
    ["test_workspace_resource_contract.py", "前端工作台资源卡片契约"],
    ["tests/application/*", "Session/Profile/Resource/Tutor/Assessment 服务"],
    ["test_auth.py", "鉴权、Token、验证码或用户流程"],
    ["test_http_observability.py", "请求 ID、指标、日志上下文"],
])
w.heading("7.4 运行建议", 2)
w.numbered([
    "本地开发前端：进入 frontend 后执行 npm install 与 npm run dev，API 走代理或同源配置。",
    "本地后端：安装 requirements 后执行 python frontend/server.py，默认服务端口以 server.py 配置为准，Docker 中暴露 8800。",
    "完整依赖：使用 docker-compose up 启动 app、Neo4j、Elasticsearch、Milvus、MinIO、etcd。",
    "测试：使用 pytest 运行 tests；对 LLM 相关测试建议使用 mock 或测试 Provider，避免真实费用和不稳定。",
    "观测：通过 X-Request-ID 定位一次请求，在 /api/ops/metrics 查看指标快照，在日志中查看 Agent 阶段事件。",
])

w.heading("8. 风险点、优化建议与迭代路线", 1)
w.heading("8.1 当前风险点", 2)
w.table(["风险", "影响", "建议"], [
    [".env.example 中出现真实格式的 DashScope Key", "密钥泄露、费用风险、供应链风险", "立即轮换密钥，提交历史中移除真实密钥，示例文件统一使用占位符"],
    ["frontend/server.py 同时承担前端目录与后端入口职责", "新成员容易误解部署入口，职责边界不清", "可迁移到 src/web/app.py 或 backend/server.py，并保留兼容启动脚本"],
    ["LLM 输出不稳定", "资源格式、Tutor 回答和诊断报告可能波动", "强化结构化输出 schema、重试、校验和回退模板"],
    ["多存储依赖较多", "本地部署门槛和故障面增加", "提供 lite profile：SQLite + 内存图谱 + mock vector，生产再启用完整依赖"],
    ["AgentState 可能变大", "SessionResponse 过大、前端刷新成本增加", "做分页/按节点懒加载、压缩 agent_feedback、资源差量更新"],
    ["SSE 与 Token 刷新边界", "长连接中 Token 过期或网络中断体验受影响", "增加心跳、断点续传、取消控制和统一错误事件"],
    ["资源渲染安全", "Markdown/Mermaid/HTML 内容可能引入 XSS 或渲染崩溃", "继续使用 DOMPurify，并在后端限制 HTML，Mermaid 渲染前做语法白名单"],
])
w.heading("8.2 架构优化建议", 2)
w.bullets([
    "将 Web 入口、静态托管、API 路由和服务装配进一步模块化，便于单元测试和未来拆分微服务。",
    "为 SessionResponse、ResourceCard、AgentFeedbackItem 建立显式版本迁移策略，并在前端处理未知字段和旧版本字段。",
    "引入任务队列处理高成本资源生成，将同步请求与异步生成分离，前端通过状态轮询或 SSE 获取结果。",
    "对知识图谱查询、Planner 结果和资源生成建立多级缓存，减少 LLM 和数据库压力。",
    "建立端到端 tracing：request_id + session_id + agent stage + provider request id，便于复现问题。",
    "补充 Playwright 端到端测试，覆盖登录、选课、画像、资源加载、Tutor 流式回答和行为提交。",
])
w.heading("8.3 迭代路线", 2)
w.table(["阶段", "目标", "交付物"], [
    ["短期", "提升稳定性与安全", "密钥轮换、示例环境清理、SSE 错误处理、核心接口契约测试"],
    ["中期", "提升个性化效果", "更细粒度画像、资源偏好学习、错题本、学习报告导出"],
    ["中期", "提升工程可维护性", "Web 入口重构、服务装配模块化、统一异常码、OpenAPI 文档"],
    ["长期", "平台化与规模化", "异步任务队列、资源市场、多课程图谱管理、班级/教师端、A/B 实验"],
])

w.heading("9. 附录：关键文件地图", 1)
w.table(["文件", "说明"], [
    ["frontend/server.py", "Starlette 入口、API 路由、静态资源托管、兼容接口"],
    ["src/orchestration_core.py", "官方学习步骤、多智能体编排、ColdStartOrchestrator、LearningStepResult"],
    ["src/orchestration_runtime.py", "运行时单例、LLM Provider 顺序、Agent 节点装配"],
    ["src/state/agent_state.py", "AgentState、StaticProfile、DynamicProfile、LatestBehavior、ResourceCard"],
    ["src/agents/evaluator_node.py", "行为评估、异常检测、PID 重规划"],
    ["src/agents/profiler_node.py", "学习者画像与掌握度更新"],
    ["src/agents/planner_node.py", "DAG-Dijkstra 路径规划"],
    ["src/agents/content_mesh_node.py", "资源调度、WFQ、影子预生成"],
    ["src/agents/validator_node.py", "资源反幻觉校验"],
    ["src/agents/tutor_node.py", "导师答疑 Agent"],
    ["src/agents/assessment_node.py", "能力雷达、诊断报告、策略控制"],
    ["src/agents/resource_generation_agent.py", "统一资源生成 Agent"],
    ["src/application/session_service.py", "会话应用服务"],
    ["src/application/profile_service.py", "画像应用服务"],
    ["src/application/resource_service.py", "资源应用服务"],
    ["src/application/tutor_service.py", "Tutor 应用服务"],
    ["src/application/assessment_service.py", "评测应用服务"],
    ["src/database/models_sqlalchemy.py", "数据库 ORM 模型"],
    ["src/auth/routes.py", "鉴权相关路由实现"],
    ["src/graph/knowledge_graph_manager.py", "知识图谱管理"],
    ["src/vector/milvus_client.py", "Milvus 向量数据库客户端"],
    ["src/vector/elasticsearch_knowledge_base.py", "Elasticsearch 知识库检索"],
    ["src/llm/client_v2.py", "LLM Provider 新版客户端"],
    ["src/observability.py", "结构化日志与指标"],
    ["frontend/src/composables/useEduAgent.js", "前端核心状态机"],
    ["frontend/src/services/apiClient.js", "Axios 客户端与 Token 刷新"],
    ["frontend/src/services/eduAgentApi.js", "前端业务 API 封装"],
    ["frontend/src/views/AppWorkspaceView.vue", "工作台状态容器"],
    ["frontend/src/components/ResourceCanvas.vue", "资源卡片画布"],
    ["frontend/src/components/ChatArea.vue", "Tutor/画像聊天区"],
    ["frontend/src/components/KnowledgeTree.vue", "知识路径与图谱展示"],
    ["Dockerfile", "多阶段构建和应用启动"],
    ["docker-compose.yml", "完整依赖编排"],
    ["requirements.txt", "Python 依赖"],
    ["frontend/package.json", "前端依赖和构建脚本"],
])
w.heading("10. 结语", 1)
w.para("从当前源码看，EduAgent 已经形成较完整的智能教育系统原型：前端具备完整学习工作台体验，后端具备多智能体学习闭环、资源生成、路径规划、鉴权、持久化、图谱和检索能力。后续最值得优先投入的是安全清理、接口契约稳定化、异步资源生成、端到端测试和运行观测体系完善。")

w.save()
print(json.dumps({"docx": str(DOCX_PATH), "md": str(MD_PATH), "paragraphs": len(w.doc.paragraphs), "tables": len(w.doc.tables)}, ensure_ascii=False, indent=2))
